#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Import one-off : Plan de Nettoyage et Désinfection (PND) du site C3ES.

Source : "/srv/ba38/uploads/cuisine/PND C3ES.xlsx" — 14 onglets, un par
zone du site (Chambres froides, Zone cuisson, Vestiaires, ...). Chaque
onglet liste les surfaces à nettoyer dans cette zone (colonnes SURFACE /
fréquence mini / produit et dose / POINT CLEF), précédées d'une ligne
"locaux" donnant la description complète de la zone (n° de salle etc.).

Remplace l'ancien registre "nettoyage" à zone unique
(cuisine_hygiene_zones_nettoyage / cuisine_hygiene_nettoyages, laissées en
base mais plus utilisées par l'écran) par ce référentiel bien plus fin :
une zone = plusieurs surfaces, chacune avec sa propre fréquence.

Petites corrections d'orthographe/espaces du fichier source normalisées
(ex. "apres" → "après", double espace supprimé) — le contenu métier n'est
pas modifié.

Idempotent (upsert par zone.code, puis par (zone_id, nom_surface)).
Utilisable sur DEV ou PROD via --env (défaut : dev).
"""

import argparse
import os
import re
import sys
import sqlite3
import unicodedata
from contextlib import closing

import openpyxl

DB_PATHS = {
    "dev": "/srv/ba38/dev/instance/ba380dev.sqlite",
    "prod": "/srv/ba38/prod/instance/ba380.sqlite",
    "dev_test": "/srv/ba38/dev/instance/ba380dev_test.sqlite",
    "prod_test": "/srv/ba38/prod/instance/ba380_test.sqlite",
}

FICHIER_SOURCE = "/srv/ba38/uploads/cuisine/PND C3ES.xlsx"

NORMALISATION = {
    "apres chaque utilisation": "après chaque utilisation",
    "ASTRASURF 1%  auto": "ASTRASURF 1% auto",
}


def _slugify(libelle: str) -> str:
    txt = unicodedata.normalize("NFKD", libelle).encode("ascii", "ignore").decode("ascii")
    txt = re.sub(r"[^a-zA-Z0-9]+", "_", txt).strip("_").lower()
    return txt


def _nettoie(valeur):
    if valeur is None:
        return None
    valeur = str(valeur).strip()
    valeur = re.sub(r"\s+", " ", valeur)
    return NORMALISATION.get(valeur, valeur) or None


def lire_pnd(fichier):
    wb = openpyxl.load_workbook(fichier, data_only=True)
    zones = []
    for ordre, nom_onglet in enumerate(wb.sheetnames, start=1):
        ws = wb[nom_onglet]
        rows = list(ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=True))

        description = None
        header_idx = None
        for i, r in enumerate(rows):
            if r[0] == "locaux":
                description = _nettoie(r[1])
            if r[0] == "SURFACE":
                header_idx = i
                break
        if header_idx is None:
            print(f"⚠️  Onglet « {nom_onglet} » : pas d'en-tête SURFACE trouvé, ignoré.")
            continue

        surfaces = []
        for r in rows[header_idx + 1:]:
            if r[0] is None or r[0] == "MÉTHODE :":
                break
            surfaces.append({
                "nom_surface": _nettoie(r[0]),
                "frequence": _nettoie(r[1]),
                "produit_dose": _nettoie(r[2]) or "aucun",
                "point_clef": _nettoie(r[3]) if len(r) > 3 else None,
            })

        libelle = nom_onglet.strip()
        zones.append({
            "code": _slugify(libelle),
            "libelle": libelle,
            "description": description,
            "ordre": ordre,
            "surfaces": surfaces,
        })
    return zones


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--env", choices=["dev", "prod", "dev_test", "prod_test"], default="dev")
    parser.add_argument("--fichier", default=FICHIER_SOURCE)
    args = parser.parse_args()

    db_path = DB_PATHS[args.env]
    if args.env == "prod" and "/prod/" not in db_path:
        raise RuntimeError(f"❌ Chemin PROD invalide : {db_path}")
    if args.env == "dev" and "/dev/" not in db_path:
        raise RuntimeError(f"❌ Chemin DEV invalide : {db_path}")

    if not os.path.exists(db_path):
        print(f"❌ Base introuvable : {db_path}")
        sys.exit(1)
    if not os.path.exists(args.fichier):
        print(f"❌ Fichier Excel introuvable : {args.fichier}")
        sys.exit(1)

    zones = lire_pnd(args.fichier)
    total_surfaces = sum(len(z["surfaces"]) for z in zones)
    print(f"➡ {len(zones)} zone(s), {total_surfaces} surface(s) lues dans {args.fichier}")

    with closing(sqlite3.connect(db_path)) as conn:
        cur = conn.cursor()
        zones_crees, surfaces_crees, surfaces_maj = 0, 0, 0

        for zone in zones:
            existante = cur.execute(
                "SELECT id FROM cuisine_pnd_zones WHERE code = ?", (zone["code"],)
            ).fetchone()
            if existante:
                zone_id = existante[0]
                cur.execute(
                    "UPDATE cuisine_pnd_zones SET libelle = ?, description = ?, ordre = ? WHERE id = ?",
                    (zone["libelle"], zone["description"], zone["ordre"], zone_id),
                )
            else:
                cur.execute(
                    "INSERT INTO cuisine_pnd_zones (code, libelle, description, ordre) VALUES (?, ?, ?, ?)",
                    (zone["code"], zone["libelle"], zone["description"], zone["ordre"]),
                )
                zone_id = cur.lastrowid
                zones_crees += 1

            for ordre_surface, surface in enumerate(zone["surfaces"], start=1):
                if not surface["nom_surface"]:
                    continue
                existante_s = cur.execute(
                    "SELECT id FROM cuisine_pnd_surfaces WHERE zone_id = ? AND nom_surface = ?",
                    (zone_id, surface["nom_surface"]),
                ).fetchone()
                if existante_s:
                    cur.execute(
                        """UPDATE cuisine_pnd_surfaces
                           SET frequence = ?, produit_dose = ?, point_clef = ?, ordre = ?
                           WHERE id = ?""",
                        (surface["frequence"], surface["produit_dose"], surface["point_clef"],
                         ordre_surface, existante_s[0]),
                    )
                    surfaces_maj += 1
                else:
                    cur.execute(
                        """INSERT INTO cuisine_pnd_surfaces
                           (zone_id, nom_surface, frequence, produit_dose, point_clef, ordre)
                           VALUES (?, ?, ?, ?, ?, ?)""",
                        (zone_id, surface["nom_surface"], surface["frequence"],
                         surface["produit_dose"], surface["point_clef"], ordre_surface),
                    )
                    surfaces_crees += 1

        conn.commit()
        total_zones = cur.execute("SELECT COUNT(*) FROM cuisine_pnd_zones").fetchone()[0]
        total_surf = cur.execute("SELECT COUNT(*) FROM cuisine_pnd_surfaces").fetchone()[0]

    print(f"✓ Zones : {zones_crees} créée(s). Surfaces : {surfaces_crees} créée(s), {surfaces_maj} mise(s) à jour.")
    print(f"📋 Total en base [{args.env}] : {total_zones} zone(s), {total_surf} surface(s).")


if __name__ == "__main__":
    main()
