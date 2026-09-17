#!/usr/bin/env python3
"""Reprise ponctuelle des deux fichiers Excel maintenus à la main pour le
suivi pluriannuel des magasins collectés (« collectes magasins evolution.xlsx »
et « liste-magasins-réferentiel.xlsx », /srv/ba38/uploads/) vers deux tables
en base : collecte_magasins_referentiel (fiche magasin, cumulative, tous les
magasins jamais collectés) et collecte_magasins_resultats (résultat en kg par
magasin et par année, une ligne par couple).

La feuille « non collecté 2025 » de collectes magasins evolution.xlsx est un
sous-ensemble de la feuille « magasins » (mêmes Code VIF) — elle n'est jamais
lue séparément.

Idempotent : n'écrase pas une valeur déjà en base pour un même magasin/année,
sauf option --forcer.
"""
import argparse
import os
import re
import sys
from datetime import datetime

from openpyxl import load_workbook

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))
from ba38_utilitaires.core import get_db_connection

FICHIER_EVOLUTION = "/srv/ba38/uploads/collectes magasins evolution.xlsx"
FICHIER_REFERENTIEL = "/srv/ba38/uploads/liste-magasins-réferentiel.xlsx"


def creer_tables(conn):
    conn.execute("""
        CREATE TABLE IF NOT EXISTS collecte_magasins_referentiel (
            code_vif TEXT PRIMARY KEY,
            nom TEXT,
            etat TEXT,
            adresse TEXT,
            ville TEXT,
            code_postal TEXT,
            telephone TEXT,
            email TEXT,
            stockage TEXT,
            gardee_par TEXT,
            ajoute_le TEXT,
            ajoute_par TEXT
        )
    """)
    conn.execute("""
        CREATE TABLE IF NOT EXISTS collecte_magasins_resultats (
            code_vif TEXT NOT NULL,
            annee INTEGER NOT NULL,
            resultat_kg REAL,
            importe_le TEXT,
            importe_par TEXT,
            PRIMARY KEY (code_vif, annee)
        )
    """)
    conn.commit()


def importer_referentiel(conn):
    wb = load_workbook(FICHIER_REFERENTIEL, data_only=True)
    ws = wb["Sheet1"]
    maintenant = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    nb = 0
    for r in range(2, ws.max_row + 1):
        code_vif = ws.cell(r, 1).value
        if not code_vif or str(code_vif).strip().lower() == "total":
            continue
        code_vif = str(code_vif).strip()
        conn.execute("""
            INSERT INTO collecte_magasins_referentiel
                (code_vif, nom, etat, adresse, ville, code_postal, telephone, email, stockage, gardee_par, ajoute_le, ajoute_par)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            ON CONFLICT(code_vif) DO NOTHING
        """, (
            code_vif,
            ws.cell(r, 2).value,
            ws.cell(r, 3).value,
            ws.cell(r, 4).value,
            ws.cell(r, 5).value,
            str(ws.cell(r, 6).value or ""),
            str(ws.cell(r, 7).value or ""),
            ws.cell(r, 8).value,
            ws.cell(r, 9).value,
            ws.cell(r, 10).value,
            maintenant,
            "import_initial",
        ))
        nb += 1
    conn.commit()
    print(f"📇 Référentiel : {nb} magasins traités")


def importer_resultats(conn, forcer=False):
    wb = load_workbook(FICHIER_EVOLUTION, data_only=True)
    ws = wb["magasins"]
    entetes = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    colonnes_annees = {}
    for i, entete in enumerate(entetes, start=1):
        if not entete:
            continue
        correspondance = re.match(r"resultats magasins (\d{4})", str(entete).strip())
        if correspondance:
            colonnes_annees[int(correspondance.group(1))] = i
    print(f"📅 Années détectées dans le fichier : {sorted(colonnes_annees)}")

    maintenant = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    nb = 0
    for r in range(2, ws.max_row + 1):
        code_vif = ws.cell(r, 1).value
        if not code_vif or str(code_vif).strip().lower() == "total":
            continue
        code_vif = str(code_vif).strip()
        for annee, col in colonnes_annees.items():
            valeur = ws.cell(r, col).value
            if valeur in (None, "", 0):
                continue
            if forcer:
                conn.execute("""
                    INSERT INTO collecte_magasins_resultats (code_vif, annee, resultat_kg, importe_le, importe_par)
                    VALUES (?, ?, ?, ?, ?)
                    ON CONFLICT(code_vif, annee) DO UPDATE SET
                        resultat_kg = excluded.resultat_kg, importe_le = excluded.importe_le, importe_par = excluded.importe_par
                """, (code_vif, annee, float(valeur), maintenant, "import_initial"))
            else:
                conn.execute("""
                    INSERT INTO collecte_magasins_resultats (code_vif, annee, resultat_kg, importe_le, importe_par)
                    VALUES (?, ?, ?, ?, ?)
                    ON CONFLICT(code_vif, annee) DO NOTHING
                """, (code_vif, annee, float(valeur), maintenant, "import_initial"))
            nb += 1
    conn.commit()
    print(f"📊 Résultats : {nb} valeurs (magasin, année) traitées")


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--forcer", action="store_true", help="Écrase les résultats déjà en base pour un même magasin/année")
    args = parser.parse_args()

    with get_db_connection() as conn:
        creer_tables(conn)
        importer_referentiel(conn)
        importer_resultats(conn, forcer=args.forcer)


if __name__ == "__main__":
    main()
