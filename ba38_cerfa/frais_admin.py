"""CERFA abandon de frais bénévoles — écrans trésorerie : campagnes,
paramètres (barème, ticket TAG, textes), lieux, population, suivi,
validation, documents, export."""

import io
import json
import os
import re
from datetime import date, datetime

from flask import abort, flash, jsonify, redirect, render_template, request, send_file, session, url_for
from flask_login import current_user, login_required
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from ba38_utilitaires.core import adresse_test_utilisateur, get_db_connection, require_access, write_log

from ba38_cerfa import cerfa_bp
from ba38_cerfa.calculs import (
    adresse_complete, bareme_campagne, champs_manquants, charger_json, geocoder, totaux_declaration,
)
from ba38_cerfa.constants import (
    BAREME_DEFAUT, BAREME_SOURCE_DEFAUT, COURRIER_REMBOURSEMENT_TEXTE, MAIL_CERFA_CORPS, MAIL_CERFA_SUJET,
    MAIL_INVITATION_CORPS, MAIL_INVITATION_SUJET, MAIL_RELANCE_CORPS, MAIL_RELANCE_SUJET,
    MAIL_REMBOURSEMENT_CORPS, MAIL_REMBOURSEMENT_SUJET, MOIS, PRIX_TICKET_TAG_DEFAUT, REPLY_TO_DEFAUT,
    SIGNATAIRE_QUALITE_DEFAUT, STATUTS, STATUTS_VALIDES, TYPES_PIECE, TYPES_VEHICULE,
)
from ba38_cerfa.frais_commun import (
    construire_document, lien_questionnaire, lieux_declaration, libelle_statut, lister_lieux,
    nom_fichier_document, nouveau_token, pieces_declaration, statut_document_remis,
)
from ba38_cerfa.pdf import fusionner_pdfs

# Paramètres recopiés de la campagne précédente lors de la création d'une
# nouvelle campagne (sinon valeurs par défaut des constantes).
PARAMETRES_RECOPIES = [
    "prix_ticket_tag", "nb_tickets_par_journee", "majoration_electrique", "bareme_json", "bareme_source",
    "expediteur_email", "reply_to_email", "signataire_nom", "signataire_qualite",
    "mail_invitation_sujet", "mail_invitation_corps", "mail_relance_sujet", "mail_relance_corps",
    "mail_cerfa_sujet", "mail_cerfa_corps", "mail_remboursement_sujet", "mail_remboursement_corps",
    "courrier_remboursement_texte",
]


def _campagne_ou_404(conn, campagne_id):
    campagne = conn.execute("SELECT * FROM cerfa_frais_campagnes WHERE id = ?", (campagne_id,)).fetchone()
    if not campagne:
        abort(404)
    return campagne


def _declaration_ou_404(conn, declaration_id):
    declaration = conn.execute("SELECT * FROM cerfa_frais_declarations WHERE id = ?", (declaration_id,)).fetchone()
    if not declaration:
        abort(404)
    return declaration


def _colonne_cotisation(conn, annee):
    colonne = f"cotisation_{int(annee)}"
    colonnes = {r["name"] for r in conn.execute("PRAGMA table_info(benevoles)")}
    return colonne if colonne in colonnes else None


def _maintenant():
    return datetime.now().isoformat(timespec="seconds")


# ============================================================================
# 📅 Campagnes
# ============================================================================
@cerfa_bp.route("/frais")
@login_required
@require_access("cerfa", "lecture")
def frais_campagnes():
    with get_db_connection() as conn:
        campagnes = conn.execute("""
            SELECT c.*,
                   (SELECT COUNT(*) FROM cerfa_frais_declarations d WHERE d.campagne_id = c.id) AS nb_declarations,
                   (SELECT COUNT(*) FROM cerfa_frais_declarations d WHERE d.campagne_id = c.id
                       AND d.statut IN ('soumis','valide','termine','courrier_envoye','rembourse')) AS nb_reponses,
                   (SELECT COUNT(*) FROM cerfa_frais_declarations d WHERE d.campagne_id = c.id
                       AND d.statut IN ('valide','termine','courrier_envoye','rembourse')) AS nb_valides
            FROM cerfa_frais_campagnes c ORDER BY c.annee_frais DESC
        """).fetchall()
    aujourd_hui = date.today()
    annee_proposee = aujourd_hui.year if aujourd_hui.month >= 9 else aujourd_hui.year - 1
    return render_template("cerfa/frais/campagnes.html", campagnes=campagnes, annee_proposee=annee_proposee)


@cerfa_bp.route("/frais/creer", methods=["POST"])
@login_required
@require_access("cerfa", "ecriture")
def frais_creer_campagne():
    try:
        annee_frais = int(request.form.get("annee_frais", ""))
    except ValueError:
        flash("❌ Année invalide", "danger")
        return redirect(url_for("cerfa.frais_campagnes"))

    with get_db_connection() as conn:
        if conn.execute("SELECT 1 FROM cerfa_frais_campagnes WHERE annee_frais = ?", (annee_frais,)).fetchone():
            flash(f"⚠️ La campagne des frais {annee_frais} existe déjà", "warning")
            return redirect(url_for("cerfa.frais_campagnes"))

        precedente = conn.execute(
            "SELECT * FROM cerfa_frais_campagnes ORDER BY annee_frais DESC LIMIT 1"
        ).fetchone()

        valeurs = {
            "prix_ticket_tag": PRIX_TICKET_TAG_DEFAUT,
            "nb_tickets_par_journee": 2,
            "majoration_electrique": 0.20,
            "bareme_json": json.dumps(BAREME_DEFAUT, ensure_ascii=False),
            "bareme_source": BAREME_SOURCE_DEFAUT,
            "expediteur_email": None,
            "reply_to_email": REPLY_TO_DEFAUT,
            "signataire_nom": "",
            "signataire_qualite": SIGNATAIRE_QUALITE_DEFAUT,
            "mail_invitation_sujet": MAIL_INVITATION_SUJET,
            "mail_invitation_corps": MAIL_INVITATION_CORPS,
            "mail_relance_sujet": MAIL_RELANCE_SUJET,
            "mail_relance_corps": MAIL_RELANCE_CORPS,
            "mail_cerfa_sujet": MAIL_CERFA_SUJET,
            "mail_cerfa_corps": MAIL_CERFA_CORPS,
            "mail_remboursement_sujet": MAIL_REMBOURSEMENT_SUJET,
            "mail_remboursement_corps": MAIL_REMBOURSEMENT_CORPS,
            "courrier_remboursement_texte": COURRIER_REMBOURSEMENT_TEXTE,
        }
        if precedente:
            for cle in PARAMETRES_RECOPIES:
                if precedente[cle] is not None:
                    valeurs[cle] = precedente[cle]

        colonnes = ["annee_frais", "annee_emission", "date_limite", "date_creation", "cree_par"] + list(valeurs)
        params = [annee_frais, annee_frais + 1, f"{annee_frais + 1}-02-28", _maintenant(), current_user.email] \
            + list(valeurs.values())
        cur = conn.execute(
            f"INSERT INTO cerfa_frais_campagnes ({', '.join(colonnes)}) VALUES ({', '.join('?' * len(colonnes))})",
            params,
        )
        conn.commit()
        campagne_id = cur.lastrowid

    write_log(f"📜 CERFA frais : campagne {annee_frais} créée par {current_user.email}")
    flash(f"✅ Campagne des frais {annee_frais} créée — vérifiez les paramètres (barème, ticket TAG, date limite).", "success")
    return redirect(url_for("cerfa.frais_parametres", campagne_id=campagne_id))


# ============================================================================
# ⚙️ Paramètres d'une campagne
# ============================================================================
CHAMPS_TEXTE = [
    "bareme_source", "expediteur_email", "reply_to_email", "signataire_nom", "signataire_qualite",
    "mail_invitation_sujet", "mail_invitation_corps", "mail_relance_sujet", "mail_relance_corps",
    "mail_cerfa_sujet", "mail_cerfa_corps", "mail_remboursement_sujet", "mail_remboursement_corps",
    "courrier_remboursement_texte",
]


def _nombre(valeur, defaut=None):
    try:
        return float(str(valeur).replace(",", ".").strip())
    except (TypeError, ValueError):
        return defaut


@cerfa_bp.route("/frais/<int:campagne_id>/parametres", methods=["GET", "POST"])
@login_required
@require_access("cerfa", "ecriture")
def frais_parametres(campagne_id):
    with get_db_connection() as conn:
        campagne = _campagne_ou_404(conn, campagne_id)
        bareme = bareme_campagne(campagne)

        if request.method == "POST":
            # Barème : champs bareme-<type>-<index>-<coef> + seuils bareme-<type>-seuil1/2
            nouveau = {}
            for type_v, categorie in bareme.items():
                seuils = [
                    _nombre(request.form.get(f"bareme-{type_v}-seuil1"), categorie["seuils"][0]),
                    _nombre(request.form.get(f"bareme-{type_v}-seuil2"), categorie["seuils"][1]),
                ]
                lignes = []
                for i, ligne in enumerate(categorie["lignes"]):
                    lignes.append({
                        "label": ligne["label"],
                        "cv_max": ligne["cv_max"],
                        **{k: _nombre(request.form.get(f"bareme-{type_v}-{i}-{k}"), ligne[k])
                           for k in ("a1", "a2", "b2", "a3")},
                    })
                nouveau[type_v] = {"seuils": seuils, "lignes": lignes}

            valeurs = {k: (request.form.get(k) or "").replace("\r\n", "\n").strip() or None for k in CHAMPS_TEXTE}
            valeurs.update({
                "annee_emission": int(_nombre(request.form.get("annee_emission"), campagne["annee_emission"])),
                "date_limite": request.form.get("date_limite") or None,
                "prix_ticket_tag": _nombre(request.form.get("prix_ticket_tag"), campagne["prix_ticket_tag"]),
                "nb_tickets_par_journee": int(_nombre(request.form.get("nb_tickets_par_journee"),
                                                      campagne["nb_tickets_par_journee"])),
                "majoration_electrique": (_nombre(request.form.get("majoration_electrique_pct"), 20) or 0) / 100,
                "bareme_json": json.dumps(nouveau, ensure_ascii=False),
            })
            conn.execute(
                f"UPDATE cerfa_frais_campagnes SET {', '.join(f'{k} = ?' for k in valeurs)} WHERE id = ?",
                list(valeurs.values()) + [campagne_id],
            )
            conn.commit()
            write_log(f"⚙️ CERFA frais : paramètres campagne {campagne['annee_frais']} modifiés par {current_user.email}")
            flash("✅ Paramètres enregistrés. Les montants des déclarations déjà validées ne changent pas "
                  "(les dévalider puis revalider pour appliquer le nouveau barème).", "success")
            return redirect(url_for("cerfa.frais_parametres", campagne_id=campagne_id))

    return render_template("cerfa/frais/parametres.html", campagne=campagne, bareme=bareme,
                           types_vehicule=TYPES_VEHICULE)


# ============================================================================
# 📍 Lieux d'activité
# ============================================================================
@cerfa_bp.route("/frais/lieux", methods=["GET", "POST"])
@login_required
@require_access("cerfa", "ecriture")
def frais_lieux():
    with get_db_connection() as conn:
        if request.method == "POST":
            lieu_id = request.form.get("lieu_id")
            nom = (request.form.get("nom") or "").strip()
            adresse = (request.form.get("adresse") or "").strip() or None
            est_autre = 1 if request.form.get("est_autre") else 0
            ordre = int(_nombre(request.form.get("ordre"), 50))
            actif = 1 if request.form.get("actif") else 0
            if not nom:
                flash("❌ Nom du lieu obligatoire", "danger")
                return redirect(url_for("cerfa.frais_lieux"))

            lat = lon = label = None
            if adresse and not est_autre:
                geo = geocoder(adresse)
                if geo:
                    lat, lon, label, score = geo
                    if score is not None and score < 0.6:
                        flash(f"⚠️ Adresse de « {nom} » reconnue avec un doute : {label} — vérifiez-la.", "warning")
                else:
                    flash(f"⚠️ Adresse de « {nom} » introuvable : les km ne pourront pas être calculés automatiquement.",
                          "warning")

            if lieu_id:
                conn.execute("""
                    UPDATE cerfa_frais_lieux SET nom = ?, adresse = ?, latitude = ?, longitude = ?,
                        geocode_label = ?, est_autre = ?, ordre = ?, actif = ? WHERE id = ?
                """, (nom, adresse, lat, lon, label, est_autre, ordre, actif, lieu_id))
            else:
                conn.execute("""
                    INSERT INTO cerfa_frais_lieux (nom, adresse, latitude, longitude, geocode_label, est_autre, ordre, actif)
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?)
                """, (nom, adresse, lat, lon, label, est_autre, ordre, actif))
            conn.commit()
            write_log(f"📍 CERFA frais : lieu « {nom} » enregistré par {current_user.email}")
            flash(f"✅ Lieu « {nom} » enregistré", "success")
            return redirect(url_for("cerfa.frais_lieux", retour=request.args.get("retour")))

        lieux = lister_lieux(conn, actifs_seulement=False)
    return render_template("cerfa/frais/lieux.html", lieux=lieux, retour=request.args.get("retour"))


# ============================================================================
# 👥 Population : bénévoles à jour de cotisation de l'année des frais
# ============================================================================
@cerfa_bp.route("/frais/<int:campagne_id>/population", methods=["POST"])
@login_required
@require_access("cerfa", "ecriture")
def frais_population(campagne_id):
    with get_db_connection() as conn:
        campagne = _campagne_ou_404(conn, campagne_id)
        colonne = _colonne_cotisation(conn, campagne["annee_frais"])
        if not colonne:
            flash(f"❌ Pas de colonne cotisation_{campagne['annee_frais']} dans la table bénévoles.", "danger")
            return redirect(url_for("cerfa.frais_suivi", campagne_id=campagne_id))

        benevoles = conn.execute(f"""
            SELECT * FROM benevoles WHERE LOWER(TRIM(COALESCE({colonne}, ''))) = 'oui'
        """).fetchall()
        existants = {r["benevole_id"]: r for r in conn.execute(
            "SELECT * FROM cerfa_frais_declarations WHERE campagne_id = ?", (campagne_id,))}

        ajoutes = maj = 0
        for b in benevoles:
            existant = existants.get(b["id"])
            if existant is None:
                _inserer_declaration(conn, campagne_id, b)
                ajoutes += 1
            elif existant["statut"] in ("a_inviter", "invite") and not existant["derniere_modif_le"]:
                # Pas encore commencé : on reprend les coordonnées à jour de la fiche bénévole
                conn.execute("""
                    UPDATE cerfa_frais_declarations SET civilite = ?, nom = ?, prenom = ?, rue = ?,
                        complement_adresse = ?, code_postal = ?, ville = ?, email = ?, telephone = ?
                    WHERE id = ?
                """, (b["civilite"], b["nom"], b["prenom"], b["rue"], b["complement_adresse"], b["code_postal"],
                      b["ville"], (b["email"] or "").strip() or None, b["telephone_portable"], existant["id"]))
                maj += 1
        conn.commit()

    write_log(f"👥 CERFA frais {campagne['annee_frais']} : population {ajoutes} ajouté(s), {maj} mis à jour")
    flash(f"✅ {len(benevoles)} bénévole(s) à jour de cotisation {campagne['annee_frais']} : "
          f"{ajoutes} ajouté(s), {maj} coordonnées actualisées.", "success")
    return redirect(url_for("cerfa.frais_suivi", campagne_id=campagne_id))


def _inserer_declaration(conn, campagne_id, b):
    conn.execute("""
        INSERT INTO cerfa_frais_declarations
            (campagne_id, benevole_id, token, statut, civilite, nom, prenom, rue, complement_adresse,
             code_postal, ville, email, telephone)
        VALUES (?, ?, ?, 'a_inviter', ?, ?, ?, ?, ?, ?, ?, ?, ?)
    """, (campagne_id, b["id"], nouveau_token(), b["civilite"], b["nom"], b["prenom"], b["rue"],
          b["complement_adresse"], b["code_postal"], b["ville"], (b["email"] or "").strip() or None,
          b["telephone_portable"]))


@cerfa_bp.route("/frais/<int:campagne_id>/ajouter_benevole", methods=["POST"])
@login_required
@require_access("cerfa", "ecriture")
def frais_ajouter_benevole(campagne_id):
    """Ajout manuel d'un bénévole hors population (ex. cotisation réglée tard)."""
    m = re.match(r"\s*(\d+)", request.form.get("benevole", ""))
    with get_db_connection() as conn:
        _campagne_ou_404(conn, campagne_id)
        b = conn.execute("SELECT * FROM benevoles WHERE id = ?", (int(m.group(1)),)).fetchone() if m else None
        if not b:
            flash("❌ Bénévole introuvable", "danger")
        elif conn.execute("SELECT 1 FROM cerfa_frais_declarations WHERE campagne_id = ? AND benevole_id = ?",
                          (campagne_id, b["id"])).fetchone():
            flash(f"ℹ️ {b['prenom']} {b['nom']} fait déjà partie de la campagne", "info")
        else:
            _inserer_declaration(conn, campagne_id, b)
            conn.commit()
            flash(f"✅ {b['prenom']} {b['nom']} ajouté(e) à la campagne", "success")
    return redirect(url_for("cerfa.frais_suivi", campagne_id=campagne_id))


# ============================================================================
# 📊 Suivi
# ============================================================================
def _lignes_suivi(conn, campagne):
    declarations = conn.execute(
        "SELECT * FROM cerfa_frais_declarations WHERE campagne_id = ? ORDER BY nom COLLATE NOCASE, prenom",
        (campagne["id"],),
    ).fetchall()
    pieces = {}
    for p in conn.execute("""
        SELECT p.declaration_id, p.type_piece FROM cerfa_frais_pieces p
        JOIN cerfa_frais_declarations d ON d.id = p.declaration_id WHERE d.campagne_id = ?
    """, (campagne["id"],)):
        pieces.setdefault(p["declaration_id"], []).append(p)
    lieux_tous = lister_lieux(conn, actifs_seulement=False)

    lignes = []
    for d in declarations:
        lieux = [l for l in lieux_tous if l["actif"] or str(l["id"]) in charger_json(d["journees_json"])]
        totaux = totaux_declaration(d, campagne, lieux)
        manquants = champs_manquants(d, campagne, lieux, pieces.get(d["id"], [])) \
            if d["statut"] in ("en_cours", "soumis") else []
        fige = d["statut"] in STATUTS_VALIDES
        lignes.append({
            "id": d["id"],
            "benevole_id": d["benevole_id"],
            "nom": d["nom"] or "",
            "prenom": d["prenom"] or "",
            "email": d["email"] or "",
            "statut": d["statut"],
            "statut_libelle": libelle_statut(d["statut"]),
            "imposable": {1: "Imposable", 0: "Non imposable"}.get(d["imposable"], ""),
            "journees": d["total_journees"] if fige else totaux["total_journees"],
            "journees_par_lieu": {str(p["lieu"]["id"]): p["journees"] for p in totaux["par_lieu"]},
            "distance": d["distance_totale"] if fige else totaux["distance_totale"],
            "montant": (d["montant_arrondi"] if d["imposable"] == 1 else d["montant"]) if fige
                       else (totaux["montant_arrondi"] if d["imposable"] == 1 else totaux["montant"]),
            "nb_pieces": len(pieces.get(d["id"], [])),
            "manquants": "; ".join(manquants),
            "numero": d["numero_document"] or "",
            "invitation": _etat_envoi(d["invitation_envoyee_le"], d["invitation_mode_test"], d["invitation_erreur"],
                                      d["invitation_statut_final"] or d["invitation_mailjet_status"]),
            "relances": d["nb_relances"] or 0,
            "document": _etat_envoi(d["document_envoye_le"], d["document_mode_test"], d["document_erreur"],
                                    d["document_statut_final"] or d["document_mailjet_status"]),
            "soumis_le": (d["soumis_le"] or "")[:10],
            "rembourse_le": (d["rembourse_le"] or "")[:10],
            "complet": d["statut"] == "soumis" and not manquants,
            "url": url_for("cerfa.frais_detail", declaration_id=d["id"]),
            "lien": lien_questionnaire(d),
        })
    return lignes


def _etat_envoi(le, mode_test, erreur, statut):
    if not le:
        return ""
    if erreur:
        return f"❌ {le[:10]} erreur"
    if mode_test:
        return f"🧪 {le[:10]} test"
    return f"✅ {le[:10]}" + (f" ({statut})" if statut else "")


def _recapitulatif(lignes, lieux):
    def compte(*statuts, situation=None):
        return [l for l in lignes if l["statut"] in statuts and (situation is None or l["imposable"] == situation)]

    r = {
        "nb": len(lignes),
        "sans_email": sum(1 for l in lignes if not l["email"]),
        "invites": sum(1 for l in lignes if l["statut"] != "a_inviter"),
        "en_cours": len(compte("en_cours")),
        "soumis": len(compte("soumis")),
        "a_envoyer": len(compte("valide")),
        "termines": len(compte("termine", "rembourse")),
        "a_rembourser": len(compte("courrier_envoye")),
        "declines": len(compte("decline")),
    }
    repondus = compte("soumis", *STATUTS_VALIDES)
    r["imposables"] = sum(1 for l in repondus if l["imposable"] == "Imposable")
    r["non_imposables"] = sum(1 for l in repondus if l["imposable"] == "Non imposable")
    r["journees_par_lieu"] = [
        (lieu["nom"], sum(l["journees_par_lieu"].get(str(lieu["id"]), 0) for l in repondus)) for lieu in lieux
    ]
    r["journees_total"] = sum(l["journees"] or 0 for l in repondus)
    r["km_total"] = round(sum(l["distance"] or 0 for l in repondus if l["imposable"] == "Imposable"), 1)
    cerfa = compte(*STATUTS_VALIDES, situation="Imposable")
    remb = compte(*STATUTS_VALIDES, situation="Non imposable")
    r["nb_cerfa"] = len(cerfa)
    r["somme_cerfa"] = sum(l["montant"] or 0 for l in cerfa)
    r["nb_remboursements"] = len(remb)
    r["somme_remboursements"] = round(sum(l["montant"] or 0 for l in remb), 2)
    r["somme_remboursee"] = round(sum(l["montant"] or 0 for l in remb if l["statut"] == "rembourse"), 2)
    r["reste_a_rembourser"] = round(r["somme_remboursements"] - r["somme_remboursee"], 2)
    return r


@cerfa_bp.route("/frais/<int:campagne_id>")
@login_required
@require_access("cerfa", "lecture")
def frais_suivi(campagne_id):
    with get_db_connection() as conn:
        campagne = _campagne_ou_404(conn, campagne_id)
        lignes = _lignes_suivi(conn, campagne)
        lieux = lister_lieux(conn, actifs_seulement=False)
        lieux_sans_adresse = [l["nom"] for l in lieux if l["actif"] and not l["est_autre"] and l["latitude"] is None]
        benevoles = conn.execute(
            "SELECT id, nom, prenom FROM benevoles ORDER BY nom COLLATE NOCASE, prenom").fetchall()
        colonne_ok = _colonne_cotisation(conn, campagne["annee_frais"]) is not None

    mail_mode = session.get("MAIL_MODE", os.getenv("MAIL_MODE", "TEST").upper())
    return render_template(
        "cerfa/frais/suivi.html",
        campagne=campagne,
        lignes=lignes,
        recap=_recapitulatif(lignes, [l for l in lieux if l["actif"]]),
        lieux_sans_adresse=lieux_sans_adresse,
        benevoles=benevoles,
        colonne_ok=colonne_ok,
        statuts=STATUTS,
        mail_mode=mail_mode,
        mail_test_to=adresse_test_utilisateur(),
        a_inviter=sum(1 for l in lignes if l["statut"] == "a_inviter" and l["email"]),
        a_relancer=sum(1 for l in lignes if l["email"] and l["statut"] in
                       (("a_inviter", "invite", "en_cours") if mail_mode == "TEST" else ("invite", "en_cours"))),
        documents_a_envoyer=sum(1 for l in lignes if l["statut"] == "valide" and l["email"]),
        completes=sum(1 for l in lignes if l["complet"]),
    )


# ============================================================================
# 🔍 Détail d'une déclaration
# ============================================================================
@cerfa_bp.route("/frais/declaration/<int:declaration_id>")
@login_required
@require_access("cerfa", "lecture")
def frais_detail(declaration_id):
    with get_db_connection() as conn:
        d = _declaration_ou_404(conn, declaration_id)
        campagne = _campagne_ou_404(conn, d["campagne_id"])
        lieux = lieux_declaration(conn, d)
        pieces = pieces_declaration(conn, d["id"])
        totaux = totaux_declaration(d, campagne, lieux)
        manquants = champs_manquants(d, campagne, lieux, pieces)
        benevole = conn.execute("SELECT * FROM benevoles WHERE id = ?", (d["benevole_id"],)).fetchone()

    return render_template(
        "cerfa/frais/detail.html",
        d=d, campagne=campagne, lieux=lieux, pieces=pieces, totaux=totaux, manquants=manquants,
        benevole=benevole, mois=MOIS, journees=charger_json(d["journees_json"]),
        km=charger_json(d["km_json"]), km_auto=charger_json(d["km_auto_json"]),
        types_vehicule=TYPES_VEHICULE, types_piece=TYPES_PIECE,
        lien=lien_questionnaire(d), statut_libelle=libelle_statut(d["statut"]), statuts_valides=STATUTS_VALIDES,
    )


def _valider(conn, d, campagne, lieux):
    """Attribue le numéro et fige les montants au barème/prix ticket du jour."""
    totaux = totaux_declaration(d, campagne, lieux)
    numero = f"{campagne['annee_emission']}-{d['benevole_id']}"
    conn.execute("""
        UPDATE cerfa_frais_declarations SET statut = 'valide', valide_le = ?, valide_par = ?,
            numero_document = ?, total_journees = ?, distance_totale = ?, montant = ?, montant_arrondi = ?
        WHERE id = ?
    """, (_maintenant(), current_user.email, numero, totaux["total_journees"], totaux["distance_totale"],
          totaux["montant"], totaux["montant_arrondi"], d["id"]))
    write_log(f"✅ CERFA frais : déclaration {d['id']} ({d['nom']} {d['prenom']}) validée n°{numero} "
              f"montant={totaux['montant']} par {current_user.email}")
    return numero


@cerfa_bp.route("/frais/declaration/<int:declaration_id>/valider", methods=["POST"])
@login_required
@require_access("cerfa", "ecriture")
def frais_valider(declaration_id):
    with get_db_connection() as conn:
        d = _declaration_ou_404(conn, declaration_id)
        campagne = _campagne_ou_404(conn, d["campagne_id"])
        lieux = lieux_declaration(conn, d)
        manquants = champs_manquants(d, campagne, lieux, pieces_declaration(conn, d["id"]))

        if d["statut"] not in ("en_cours", "soumis"):
            flash("ℹ️ Déclaration déjà validée ou non saisie", "info")
        elif d["imposable"] not in (0, 1):
            flash("❌ Situation fiscale non renseignée : impossible de valider", "danger")
        elif manquants and not request.form.get("forcer"):
            flash("⚠️ Éléments manquants : " + "; ".join(manquants)
                  + ". Cochez « valider malgré tout » pour forcer.", "warning")
        else:
            numero = _valider(conn, d, campagne, lieux)
            conn.commit()
            flash(f"✅ Déclaration validée — n° {numero}", "success")
    return redirect(url_for("cerfa.frais_detail", declaration_id=declaration_id))


@cerfa_bp.route("/frais/<int:campagne_id>/valider_completes", methods=["POST"])
@login_required
@require_access("cerfa", "ecriture")
def frais_valider_completes(campagne_id):
    """Validation groupée : uniquement les déclarations transmises sans aucun
    élément manquant ; les autres restent à contrôler une par une."""
    nb = 0
    with get_db_connection() as conn:
        campagne = _campagne_ou_404(conn, campagne_id)
        for d in conn.execute("SELECT * FROM cerfa_frais_declarations WHERE campagne_id = ? AND statut = 'soumis'",
                              (campagne_id,)).fetchall():
            lieux = lieux_declaration(conn, d)
            if not champs_manquants(d, campagne, lieux, pieces_declaration(conn, d["id"])):
                _valider(conn, d, campagne, lieux)
                nb += 1
        conn.commit()
        restantes = conn.execute("SELECT COUNT(*) FROM cerfa_frais_declarations WHERE campagne_id = ? "
                                 "AND statut = 'soumis'", (campagne_id,)).fetchone()[0]
    write_log(f"✅ CERFA frais {campagne['annee_frais']} : validation groupée de {nb} déclaration(s) "
              f"par {current_user.email}")
    flash(f"✅ {nb} déclaration(s) complète(s) validée(s)"
          + (f" — {restantes} transmise(s) incomplète(s) à contrôler une par une." if restantes else "."),
          "success" if not restantes else "warning")
    return redirect(url_for("cerfa.frais_suivi", campagne_id=campagne_id))


@cerfa_bp.route("/frais/declaration/<int:declaration_id>/statut", methods=["POST"])
@login_required
@require_access("cerfa", "ecriture")
def frais_changer_statut(declaration_id):
    """Dévalider, rouvrir la saisie, ne participe pas, document remis à la
    main (sans email), remboursé / annuler le remboursement."""
    action = request.form.get("action")
    maintenant = _maintenant()
    with get_db_connection() as conn:
        d = _declaration_ou_404(conn, declaration_id)
        statut = d["statut"]
        if action == "devalider" and statut in ("valide", "termine", "courrier_envoye"):
            remis = (d["document_envoye_le"] and not d["document_mode_test"]) or d["document_remis_le"]
            if remis and not request.form.get("confirmer"):
                flash("⚠️ Le document a déjà été remis au bénévole : confirmez la dévalidation.", "warning")
                return redirect(url_for("cerfa.frais_detail", declaration_id=declaration_id))
            conn.execute("UPDATE cerfa_frais_declarations SET statut = 'soumis', valide_le = NULL, valide_par = NULL, "
                         "document_remis_le = NULL, document_remis_par = NULL WHERE id = ?", (d["id"],))
        elif action == "rouvrir" and statut in ("soumis", "decline"):
            conn.execute("UPDATE cerfa_frais_declarations SET statut = 'en_cours', decline_le = NULL WHERE id = ?",
                         (d["id"],))
        elif action == "decliner" and statut not in STATUTS_VALIDES:
            conn.execute("UPDATE cerfa_frais_declarations SET statut = 'decline', decline_le = ? WHERE id = ?",
                         (maintenant, d["id"]))
        elif action == "remis" and statut == "valide":
            conn.execute("UPDATE cerfa_frais_declarations SET statut = ?, document_remis_le = ?, document_remis_par = ? "
                         "WHERE id = ?", (statut_document_remis(d), maintenant, current_user.email, d["id"]))
        elif action == "rembourse" and d["imposable"] == 0 and statut in ("valide", "courrier_envoye"):
            date_remb = request.form.get("date_remboursement") or maintenant[:10]
            conn.execute("UPDATE cerfa_frais_declarations SET statut = 'rembourse', rembourse_le = ?, rembourse_par = ? "
                         "WHERE id = ?", (date_remb, current_user.email, d["id"]))
        elif action == "annuler_remboursement" and statut == "rembourse":
            remis = (d["document_envoye_le"] and not d["document_mode_test"] and not d["document_erreur"]) \
                or d["document_remis_le"]
            conn.execute("UPDATE cerfa_frais_declarations SET statut = ?, rembourse_le = NULL, rembourse_par = NULL "
                         "WHERE id = ?", ("courrier_envoye" if remis else "valide", d["id"]))
        else:
            flash("❌ Action impossible dans l'état actuel", "danger")
            return redirect(url_for("cerfa.frais_detail", declaration_id=declaration_id))
        conn.commit()
    write_log(f"🔁 CERFA frais : déclaration {declaration_id} action={action} par {current_user.email}")
    flash("✅ Statut mis à jour", "success")
    retour = request.form.get("retour") or ""
    if not retour.startswith("/") or retour.startswith("//"):
        retour = url_for("cerfa.frais_detail", declaration_id=declaration_id)
    return redirect(retour)


@cerfa_bp.route("/frais/declaration/<int:declaration_id>/commentaire", methods=["POST"])
@login_required
@require_access("cerfa", "ecriture")
def frais_commentaire(declaration_id):
    with get_db_connection() as conn:
        _declaration_ou_404(conn, declaration_id)
        conn.execute("UPDATE cerfa_frais_declarations SET commentaire_tresorerie = ? WHERE id = ?",
                     ((request.form.get("commentaire_tresorerie") or "").strip() or None, declaration_id))
        conn.commit()
    flash("✅ Commentaire enregistré", "success")
    return redirect(url_for("cerfa.frais_detail", declaration_id=declaration_id))


@cerfa_bp.route("/frais/declaration/<int:declaration_id>/piece/<int:piece_id>")
@login_required
@require_access("cerfa", "lecture")
def frais_piece(declaration_id, piece_id):
    with get_db_connection() as conn:
        piece = conn.execute("SELECT * FROM cerfa_frais_pieces WHERE id = ? AND declaration_id = ?",
                             (piece_id, declaration_id)).fetchone()
    if not piece or not os.path.exists(piece["chemin"]):
        abort(404)
    return send_file(piece["chemin"], download_name=piece["nom_original"] or os.path.basename(piece["chemin"]))


# ============================================================================
# 📄 Documents (régénérés à la demande)
# ============================================================================
@cerfa_bp.route("/frais/declaration/<int:declaration_id>/document.pdf")
@login_required
@require_access("cerfa", "lecture")
def frais_document(declaration_id):
    with get_db_connection() as conn:
        d = _declaration_ou_404(conn, declaration_id)
        campagne = _campagne_ou_404(conn, d["campagne_id"])
        try:
            contenu = construire_document(conn, d, campagne)
        except ValueError as e:
            flash(f"❌ {e}", "danger")
            return redirect(url_for("cerfa.frais_detail", declaration_id=declaration_id))
    return send_file(io.BytesIO(contenu), mimetype="application/pdf", download_name=nom_fichier_document(d))


@cerfa_bp.route("/frais/<int:campagne_id>/documents.pdf")
@login_required
@require_access("cerfa", "lecture")
def frais_documents_lot(campagne_id):
    """Tous les documents validés d'un type, en un seul PDF à imprimer."""
    type_doc = request.args.get("type", "cerfa")
    imposable = 1 if type_doc == "cerfa" else 0
    with get_db_connection() as conn:
        campagne = _campagne_ou_404(conn, campagne_id)
        declarations = conn.execute("""
            SELECT * FROM cerfa_frais_declarations
            WHERE campagne_id = ? AND statut IN ('valide','termine','courrier_envoye','rembourse') AND imposable = ?
            ORDER BY nom COLLATE NOCASE, prenom
        """, (campagne_id, imposable)).fetchall()
        if not declarations:
            flash("ℹ️ Aucun document validé de ce type", "info")
            return redirect(url_for("cerfa.frais_suivi", campagne_id=campagne_id))
        contenu = fusionner_pdfs([construire_document(conn, d, campagne) for d in declarations])
    nom = "CERFA" if imposable else "Remboursements_trajets"
    return send_file(io.BytesIO(contenu), mimetype="application/pdf",
                     download_name=f"{nom}_frais_{campagne['annee_frais']}.pdf")


# ============================================================================
# 📊 Export Excel
# ============================================================================
@cerfa_bp.route("/frais/<int:campagne_id>/export.xlsx")
@login_required
@require_access("cerfa", "lecture")
def frais_export(campagne_id):
    with get_db_connection() as conn:
        campagne = _campagne_ou_404(conn, campagne_id)
        lieux = lister_lieux(conn, actifs_seulement=False)
        declarations = conn.execute(
            "SELECT * FROM cerfa_frais_declarations WHERE campagne_id = ? ORDER BY nom COLLATE NOCASE, prenom",
            (campagne_id,)).fetchall()
        nb_pieces = {r["declaration_id"]: r["n"] for r in conn.execute("""
            SELECT declaration_id, COUNT(*) AS n FROM cerfa_frais_pieces GROUP BY declaration_id""")}
        lignes_suivi = _lignes_suivi(conn, campagne)

    wb = Workbook()
    ws = wb.active
    ws.title = "Déclarations"
    entetes = ["Id bénévole", "Civilité", "Nom", "Prénom", "Adresse", "Code postal", "Ville", "Email", "Statut",
               "Situation", "Véhicule", "Marque", "Immatriculation", "CV", "Électrique"]
    entetes += [f"Journées {l['nom']}" for l in lieux] + [f"Km A/R {l['nom']}" for l in lieux]
    entetes += ["Commentaire Autres", "Total journées", "Distance totale (km)", "Montant", "Montant arrondi",
                "N° document", "Validé le", "Remboursé le", "Pièces", "Transmis le", "Invitation", "Relances", "Document envoyé",
                "Commentaire trésorerie"]
    ws.append(entetes)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDDDDD")

    for d in declarations:
        totaux = totaux_declaration(d, campagne, lieux)
        journees = {str(p["lieu"]["id"]): p["journees"] for p in totaux["par_lieu"]}
        km = charger_json(d["km_json"])
        valide = d["statut"] in STATUTS_VALIDES
        ws.append([
            d["benevole_id"], d["civilite"], d["nom"], d["prenom"],
            " ".join(filter(None, [d["rue"], d["complement_adresse"]])), d["code_postal"], d["ville"], d["email"],
            libelle_statut(d["statut"]), {1: "Imposable", 0: "Non imposable"}.get(d["imposable"], ""),
            TYPES_VEHICULE.get(d["vehicule_type"], "") if d["imposable"] == 1 else "",
            d["vehicule_marque"], d["vehicule_immatriculation"], d["vehicule_cv"],
            "oui" if d["vehicule_electrique"] else "",
            *[journees.get(str(l["id"]), 0) for l in lieux],
            *[km.get(str(l["id"])) for l in lieux],
            d["commentaire_autres"],
            d["total_journees"] if valide else totaux["total_journees"],
            d["distance_totale"] if valide else totaux["distance_totale"],
            d["montant"] if valide else totaux["montant"],
            d["montant_arrondi"] if valide else totaux["montant_arrondi"],
            d["numero_document"], (d["valide_le"] or "")[:10], d["rembourse_le"], nb_pieces.get(d["id"], 0), (d["soumis_le"] or "")[:10],
            (d["invitation_envoyee_le"] or "")[:10], d["nb_relances"] or 0, (d["document_envoye_le"] or "")[:10],
            d["commentaire_tresorerie"],
        ])

    recap = _recapitulatif(lignes_suivi, [l for l in lieux if l["actif"]])
    ws2 = wb.create_sheet("Récapitulatif")
    for libelle, valeur in [
        ("Bénévoles de la campagne", recap["nb"]), ("dont sans email", recap["sans_email"]),
        ("Invités", recap["invites"]), ("En cours de saisie", recap["en_cours"]),
        ("Déclarations transmises (à valider)", recap["soumis"]),
        ("Validées, document à envoyer", recap["a_envoyer"]),
        ("Courriers envoyés, à rembourser", recap["a_rembourser"]), ("Terminés", recap["termines"]),
        ("Ne participent pas", recap["declines"]), ("Réponses imposables", recap["imposables"]),
        ("Réponses non imposables", recap["non_imposables"]),
        *[(f"Journées {nom}", nb) for nom, nb in recap["journees_par_lieu"]],
        ("Journées totales", recap["journees_total"]), ("Km parcourus (imposables)", recap["km_total"]),
        ("Nombre de CERFA", recap["nb_cerfa"]), ("Somme des CERFA (€)", recap["somme_cerfa"]),
        ("Nombre de remboursements", recap["nb_remboursements"]),
        ("Somme des remboursements (€)", recap["somme_remboursements"]),
        ("dont déjà remboursé (€)", recap["somme_remboursee"]),
        ("Reste à rembourser (€)", recap["reste_a_rembourser"]),
    ]:
        ws2.append([libelle, valeur])
    ws2.column_dimensions["A"].width = 40

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, download_name=f"CERFA_frais_benevoles_{campagne['annee_frais']}.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
